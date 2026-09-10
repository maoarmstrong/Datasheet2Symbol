import copy, io, json, os
from collections import Counter
from pathlib import Path
os.environ['D2S_DATA_DIR']=str(Path(__file__).parent.parent/'data'/'test')
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend.main import app
from backend import store, provider
from backend.models import Pin
import pytest

client=TestClient(app)

def compact_payload(pins):
 evidence=[];table_rows=[];compact=[]
 for pin in pins:
  evidence_ids=[]
  for item in pin['evidence']:
   if item not in evidence: evidence.append(item)
   evidence_ids.append(evidence.index(item))
  table_row_id=None
  if pin.get('table_row'):
   if pin['table_row'] not in table_rows: table_rows.append(pin['table_row'])
   table_row_id=table_rows.index(pin['table_row'])
  compact.append({key:pin[key] for key in ['number','original_name','electrical_type','part','position','kind','fact','observation','issues']} | {'evidence_ids':evidence_ids,'table_row_id':table_row_id})
 return {'evidence':evidence,'table_rows':table_rows,'pins':compact,'notes':[]}

def full_payload(pins):
 return {'pins':pins,'notes':[]}

@pytest.mark.parametrize('sample',['power','digital','connector','multipage'])
def test_roundtrip_export_and_parts(sample):
 p=client.post('/api/demo/'+sample).json()
 assert p['demo'] is True
 assert client.get('/api/projects/'+p['id']+'/pdf').status_code==200
 original={k:r['raw']['original_name'] for k,r in p['pins'].items()}
 numbers=list(p['pins'])
 p=client.patch(f"/api/projects/{p['id']}/pins",json={'revision':p['revision'],'numbers':numbers,'values':{'part':'A','review':'confirmed'}}).json()
 assert all(r['review']=='confirmed' for r in p['pins'].values())
 raw=client.get(f"/api/projects/{p['id']}/export/xlsx").content
 wb=load_workbook(io.BytesIO(raw)); rows=list(wb['Pins'].values)
 assert len(rows)==1+sum(r['raw']['kind']!='mechanical' for r in p['pins'].values())
 assert all(row[3:5]==('1','Line') for row in rows[1:])
 assert all(row[5] in ('',None) for row in rows[1:])
 assert all(isinstance(row[0],str) and row[7]=='A' for row in rows[1:])
 assert {row[0]:row[8] for row in rows[1:]}=={k:v for k,v in original.items() if p['pins'][k]['raw']['kind']!='mechanical'}
 payload=client.get(f"/api/projects/{p['id']}/export/json").content
 restored=client.post('/api/import',files={'file':('project.json',payload,'application/json')}).json()
 assert restored['id']!=p['id'] and restored['pins']==p['pins'] and restored['target']==p['target']

def test_reparse_does_not_overwrite_human_values():
 p=client.post('/api/demo/power').json();n='1'
 p=client.patch(f"/api/projects/{p['id']}/pins",json={'revision':p['revision'],'numbers':[n],'values':{'symbol_name':'VIN_USER','review':'confirmed'}}).json()
 changed=copy.deepcopy(p['pins'][n]['raw']);changed['symbol_name']='VIN_MODEL';changed['original_name']='VIN_NEW'
 store.merge(p,[changed]);store.save(p)
 assert store.effective(p['pins'][n])['symbol_name']=='VIN_USER'
 assert p['pins'][n]['review']=='conflict'
 assert len(p['pins'][n]['versions'])==2
 assert p['pins']['2']['raw']['original_name']=='VIN'
 assert not p['undo']

def test_duplicate_non_power_names_are_numbered_but_power_names_repeat():
 p=client.post('/api/demo/power').json()
 fixture=json.loads(Path('fixtures/power.json').read_text('utf-8'))
 extra=[]
 for number in ['9','10']:
  pin=copy.deepcopy(next(x for x in fixture['pins'] if x['original_name']=='NC'))
  pin['number']=number
  pin['table_row']['columns'][pin['table_row']['selected_column']]=[number]
  extra.append(pin)
 store.merge(p,extra)
 nc=sorted((store.effective(r) for r in p['pins'].values() if r['raw']['original_name']=='NC'),key=lambda pin:int(pin['number']))
 assert [pin['symbol_name'] for pin in nc]==['NC1','NC2','NC3']
 assert all(pin['original_name']=='NC' for pin in nc)
 assert [store.effective(p['pins'][n])['symbol_name'] for n in ['1','2']]==['VIN','VIN']

 p=client.post('/api/demo/digital').json()
 p=client.patch(f"/api/projects/{p['id']}/pins",json={'revision':p['revision'],'numbers':['02','03'],'values':{'symbol_name':'DUP'}}).json()
 errors=client.get(f"/api/projects/{p['id']}/validation").json()['errors']
 assert any('仅 Power 类型允许重名' in error and '02, 03' in error for error in errors)

def test_revision_and_undo():
 p=client.post('/api/demo/digital').json();base=copy.deepcopy(p)
 changed=client.patch(f"/api/projects/{p['id']}/pins",json={'revision':p['revision'],'numbers':['02'],'values':{'symbol_name':'PA0/ALT'}}).json()
 assert client.patch(f"/api/projects/{p['id']}/pins",json={'revision':p['revision'],'numbers':['02'],'values':{'symbol_name':'stale'}}).status_code==409
 undo=client.post(f"/api/projects/{p['id']}/undo",json={'revision':changed['revision']}).json()
 assert undo['pins']==base['pins']

def test_number_strict_and_formula_safe():
 p=client.post('/api/demo/digital').json()
 bad=copy.deepcopy(p['pins']['01']['raw']);bad['number']=1
 with pytest.raises(Exception): Pin.model_validate(bad)
 client.patch(f"/api/projects/{p['id']}/pins",json={'revision':p['revision'],'numbers':['01'],'values':{'symbol_name':'=1+2'}})
 wb=load_workbook(io.BytesIO(client.get(f"/api/projects/{p['id']}/export/xlsx").content))
 assert wb['Pins']['A2'].value=='01' and wb['Pins']['B2'].data_type=='s'
 assert client.get(f"/api/projects/{p['id']}/export/tsv").status_code==400

def test_pdf_and_origin_validation():
 assert client.post('/api/projects',files={'file':('bad.pdf',b'not pdf')}).status_code==400
 assert client.post('/api/demo/power',headers={'Origin':'https://evil.example'}).status_code==403

def test_manual_number_correction_and_addition():
 p=client.post('/api/demo/connector').json()
 updated=client.patch(f"/api/projects/{p['id']}/pins",json={'revision':p['revision'],'numbers':['A1'],'values':{'number':'01'}})
 assert updated.status_code==200
 p=updated.json()
 assert p['pins']['A1']['raw']['number']=='A1' and store.effective(p['pins']['A1'])['number']=='01'
 assert client.patch(f"/api/projects/{p['id']}/pins",json={'revision':p['revision'],'numbers':['A1'],'values':{'number':'A2'}}).status_code==400
 p=client.post(f"/api/projects/{p['id']}/pins",json={'revision':p['revision'],'number':'SH2','page':2}).json()
 assert any(store.effective(r)['number']=='SH2' for r in p['pins'].values())
 payload=client.get(f"/api/projects/{p['id']}/export/json").content
 reopened=client.post('/api/import',files={'file':('project.json',payload)}).json()
 assert reopened['pins']==p['pins']

def test_provider_payload_and_truncation(monkeypatch):
 fixture=json.loads(Path('fixtures/power.json').read_text('utf-8'))
 monkeypatch.setenv('DASHSCOPE_API_KEY','test-fake-key');monkeypatch.setenv('DASHSCOPE_BASE_URL','https://example.invalid/v1');monkeypatch.delenv('DASHSCOPE_MODEL',raising=False)
 seen=[]
 class FakeClient:
  def __init__(self,**kwargs): pass
  def __enter__(self): return self
  def __exit__(self,*args): pass
  def post(self,url,headers,json):
   seen.append(json)
   class Response:
    status_code=200
    def json(self): return {'choices':[{'finish_reason':'stop','message':{'content':__import__('json').dumps(full_payload(fixture['pins']))}}],'usage':{'total_tokens':10}}
   return Response()
 monkeypatch.setattr(provider.httpx,'Client',FakeClient)
 result,audit=provider.invoke(Path('fixtures/power.pdf'),[1,2],'extract',fixture['target'])
 assert len(result['pins'])==9 and audit['usage']['total_tokens']==10
 assert seen[0]['model']=='qwen3-vl-plus' and seen[0]['enable_thinking'] is False and seen[0]['response_format']=={'type':'json_object'}
 assert all(pin['group']=='' for pin in result['pins'])

def test_legacy_vl_uses_non_thinking_mode(monkeypatch):
 monkeypatch.setenv('DASHSCOPE_API_KEY','test-fake-key');monkeypatch.setenv('DASHSCOPE_BASE_URL','https://example.invalid/v1');monkeypatch.setenv('DASHSCOPE_MODEL','qwen3-vl-plus')
 fixture=json.loads(Path('fixtures/power.json').read_text('utf-8'));seen=[]
 class FakeClient:
  def __init__(self,**kwargs): pass
  def __enter__(self): return self
  def __exit__(self,*args): pass
  def post(self,url,headers,json):
   seen.append(json)
   return type('Response',(),{'status_code':200,'json':lambda self:{'choices':[{'finish_reason':'stop','message':{'content':__import__('json').dumps(full_payload(fixture['pins']))}}]}})()
 monkeypatch.setattr(provider.httpx,'Client',FakeClient)
 provider.invoke(Path('fixtures/power.pdf'),[1,2],'extract',fixture['target'])
 assert seen[0]['enable_thinking'] is False
 assert len([x for x in seen[0]['messages'][1]['content'] if x['type']=='image_url'])==2

def test_xilinx_ascii_pinout_generates_orcad_sections_and_visible_pins():
 rows = [
  'A1   IO_L1P_T0L_N0_64                  0L                 64    HP        NA',
  'A2   DONE_0                            NA                 0     CONFIG    NA',
  'A3   VCCINT                            NA                 NA    NA        NA',
  'A4   CTRL                              NA                 0     CONFIG    NA',
  'A5   CTRL                              NA                 0     CONFIG    NA',
 ] + [f'G{i:03d}   GND                               NA                 NA    NA        NA' for i in range(1, 102)]
 report = '\n'.join([
  '--  Device   : xc_test_package',
  'Pin   Pin Name                            Memory Byte Group  Bank  I/O Type  Super Logic Region',
  *rows,
 ])
 response = client.post('/api/orcad/ascii/convert', files={'file': ('pinout.txt', report.encode(), 'text/plain')}, data={'max_pins_per_section': '100', 'use_ai': 'false'})
 assert response.status_code == 200
 wb = load_workbook(io.BytesIO(response.content))
 pasted = list(wb['Paste_To_OrCAD'].values)
 assert pasted[0] == ('Number','Name','Type','Pin Visibility','Shape','PinGroup','Position','Section')
 assert len(pasted) == 1 + len(rows)
 assert all(row[3] == '1' and row[4] == 'Line' and row[5] in ('', None) for row in pasted[1:])
 section_counts = Counter(row[7] for row in pasted[1:])
 assert max(section_counts.values()) <= 100
 gnd_sections = {row[7] for row in pasted[1:] if row[1] == 'GND'}
 assert len(gnd_sections) == 2
 assert {row[0]:row[1] for row in pasted[1:] if row[0] in ('A4','A5')} == {'A4':'CTRL1','A5':'CTRL2'}
 audit=list(wb['Source_Audit'].values)
 assert audit[0][-1]=='Original Name' and {row[0]:row[-1] for row in audit[1:] if row[0] in ('A4','A5')}=={'A4':'CTRL','A5':'CTRL'}
 assert wb['Section_Summary']['B6'].value == len(section_counts)
