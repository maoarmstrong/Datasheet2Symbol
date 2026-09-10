import csv, io, json
from openpyxl import Workbook
from .store import effective

COLUMNS = ['Number','Name','Type','Pin Visibility','Shape','Function Group','Position','Part','Original Name','Kind','Review','PDF Pages','Printed Pages','Observation','Issues']
# This adapter is provisional until a version-specific Capture round-trip is recorded.
TYPES = {x:x for x in ['Input','Output','Bidirectional','Passive','Power','Open Collector','Open Emitter','3-State']}

def validate(p):
    errors = []
    if not p['target']:
        errors.append('尚未确认目标型号和封装')
    if not p['pins']:
        errors.append('没有引脚数据')
    if not any(j['stage']=='extract' and j['status']=='complete' for j in p['jobs']):
        errors.append('没有完整成功的提取任务')
    if any(j['status'] in ['running','failed'] for j in p['jobs']):
        errors.append('存在运行中或失败任务，请完成重试')
    seen = set()
    names = {}
    electrical = 0
    for r in p['pins'].values():
        pin = effective(r)
        if r['review'] != 'confirmed': errors.append(f"{pin['number']}: 未审核确认")
        if pin['number'] in seen: errors.append(f"{pin['number']}: 重复脚号")
        seen.add(pin['number'])
        if pin['kind']=='mechanical': continue
        electrical += 1
        if pin['kind']=='uncertain': errors.append(f"{pin['number']}: 结构类型未确认")
        if pin['part'] not in p['parts']: errors.append(f"{pin['number']}: Part 未分配或不存在")
        if pin['electrical_type'] not in TYPES: errors.append(f"{pin['number']}: 电气类型未确认")
        if not pin['symbol_name']: errors.append(f"{pin['number']}: Symbol 名称为空")
        else: names.setdefault(pin['symbol_name'].casefold(), []).append(pin)
    for duplicate in names.values():
        if len(duplicate) > 1 and any(pin['electrical_type'] != 'Power' for pin in duplicate):
            errors.append(f"Symbol 名称 {duplicate[0]['symbol_name']} 重复；仅 Power 类型允许重名（引脚 {', '.join(pin['number'] for pin in duplicate)}）")
    count = (p['target'] or {}).get('pin_count')
    if count and electrical != count:
        errors.append(f'目标针数 {count} 与当前非机械记录数 {electrical} 不同；确认裸露焊盘/屏蔽是否计入并修正目标针数')
    return errors

def rows(p):
    return [[q['number'],q['symbol_name'],q['electrical_type'],'1','Line','',q['position'],q['part'],q['original_name'],q['kind'],r['review'],','.join(str(e['page']) for e in q['evidence']),','.join(e['printed_page'] for e in q['evidence']),q['observation'],'; '.join(q['issues']+r['conflicts'])] for r in p['pins'].values() if (q:=effective(r))['kind']!='mechanical']

def export(p, kind):
    data = rows(p)
    if kind=='tsv':
        s=io.StringIO(newline='')
        # Header-free paste area. Reject spreadsheet formula controls rather than silently changing pin names.
        if any(str(v).startswith(('=','+','-','@','\t','\r')) or '\n' in str(v) or '\t' in str(v) for row in data for v in row):
            raise ValueError('TSV 包含可能被表格软件解释为公式或分隔符的字段，请使用文本类型 XLSX')
        csv.writer(s, delimiter='\t', lineterminator='\r\n').writerows(data)
        return s.getvalue().encode('utf-8-sig')
    wb=Workbook(); ws=wb.active; ws.title='Pins'
    ws.append(COLUMNS)
    for row in data: ws.append(row)
    for row in ws:
        for cell in row: cell.data_type='s'; cell.number_format='@'
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in 'ABCDEFGH': ws.column_dimensions[col].width=24
    ref=wb.create_sheet('Reference')
    ref.append(['Format','Generic review workbook; OrCAD adaptation deferred. Unknown/unreviewed data is preserved.'])
    ref.append(['Validation',json.dumps(validate(p),ensure_ascii=False)])
    ref.append(['Target',json.dumps(p['target'],ensure_ascii=False)])
    ref.append(['Part mapping',json.dumps({x:i+1 for i,x in enumerate(p['parts'])})])
    ref.append(['Demo',str(p['demo'])])
    for number,r in p['pins'].items(): ref.append([number,json.dumps(r,ensure_ascii=False)])
    history=wb.create_sheet('History')
    history.append(['Edit history'])
    for entry in p['history']: history.append([json.dumps(entry,ensure_ascii=False)])
    for row in history:
        for cell in row: cell.data_type='s'
    for row in ref:
        for cell in row: cell.data_type='s'
    out=io.BytesIO(); wb.save(out); return out.getvalue()
