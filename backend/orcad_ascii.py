"""Deterministic AMD/Xilinx ASCII Pinout to OrCAD spreadsheet conversion.

The vendor report remains the source of truth.  AI, when explicitly enabled,
may only refine the conservative classification of otherwise ambiguous pins.
"""
from __future__ import annotations

import io
import re
from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ORCAD_HEADERS = ['Number', 'Name', 'Type', 'Pin Visibility', 'Shape', 'PinGroup', 'Position', 'Section']
VALID_TYPES = {'Input', 'Output', 'Bidirectional', 'Passive', 'Power', 'Open Collector', 'Open Emitter', '3-State'}
VALID_POSITIONS = {'Left', 'Right', 'Top', 'Bottom'}


@dataclass
class SourcePin:
    number: str
    name: str
    memory_byte_group: str
    bank: str
    io_type: str


@dataclass
class OrCadPin:
    source: SourcePin
    electrical_type: str
    position: str
    group: str
    section: str = ''
    classification_source: str = 'Rule'
    symbol_name: str = ''


def _ground(name: str) -> bool:
    return bool(re.search(r'(^|_)(GND|VSS)(_|$)|RSVDGND|GNDADC', name))


def _power(name: str) -> bool:
    return bool(re.match(r'^(VCC|VDD|AVCC|VCCAUX|VCCO|VBATT)|^VREF_', name))


def _is_config(pin: SourcePin) -> bool:
    return pin.io_type == 'CONFIG'


def _type(pin: SourcePin) -> str:
    if _ground(pin.name) or _power(pin.name):
        return 'Power'
    if re.match(r'^(DONE|TDO)_', pin.name):
        return 'Output'
    if pin.name.startswith('IO_'):
        return 'Bidirectional'
    return 'Input'


def _position(pin: SourcePin, electrical_type: str) -> str:
    if _ground(pin.name):
        return 'Bottom'
    if _power(pin.name):
        return 'Top'
    return 'Right' if electrical_type == 'Output' else 'Left'


def _group(pin: SourcePin) -> str:
    if _ground(pin.name):
        return 'GROUND'
    if _power(pin.name):
        return 'POWER'
    if _is_config(pin):
        return 'CONFIG'
    if pin.bank != 'NA':
        return f'BANK_{pin.bank}'
    return 'CONTROL_ANALOG'


def _section_name(index: int) -> str:
    name = ''
    while True:
        name = chr(65 + index % 26) + name
        index = index // 26 - 1
        if index < 0:
            return name


def _group_sort_key(group: str) -> tuple[int, str]:
    if group == 'CONFIG':
        return 0, group
    match = re.fullmatch(r'BANK_(\d+)', group)
    if match:
        return 10 + int(match.group(1)), group
    if group == 'CONTROL_ANALOG':
        return 900, group
    if group == 'POWER':
        return 910, group
    if group == 'GROUND':
        return 920, group
    return 999, group


def parse_ascii_pinout(raw: bytes) -> tuple[str, list[SourcePin]]:
    """Parse the stable six-column table emitted by AMD/Xilinx Pinout reports."""
    text = ''
    for encoding in ('utf-8-sig', 'utf-16', 'latin-1'):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        raise ValueError('无法读取 Pinout 文本；请上传 AMD/Xilinx 导出的 ASCII Pinout 文件')
    model = re.search(r'^--\s+Device\s+:\s+(\S+)', text, flags=re.MULTILINE)
    device = (model.group(1) if model else 'XILINX_FPGA').upper()
    lines = text.splitlines()
    start = next((i for i, line in enumerate(lines) if re.match(r'^Pin\s+Pin Name\s+Memory Byte Group\s+Bank\s+I/O Type', line)), -1)
    if start < 0:
        raise ValueError('未找到 ASCII Pinout 表头；该工具当前支持 AMD/Xilinx 的标准 ASCII Pinout 格式')
    pins: list[SourcePin] = []
    for line in lines[start + 1:]:
        if not re.match(r'^\S+\s{2,}\S+', line):
            continue
        fields = re.split(r'\s{2,}', line.strip())
        if len(fields) < 5:
            continue
        pins.append(SourcePin(*fields[:5]))
    if not pins:
        raise ValueError('Pinout 表中没有可用的引脚记录')
    numbers = [pin.number for pin in pins]
    if len(numbers) != len(set(numbers)):
        raise ValueError('官方文件包含重复引脚号，已拒绝生成，避免创建错误符号')
    return device, pins


def make_orcad_pins(pins: list[SourcePin]) -> list[OrCadPin]:
    return [OrCadPin(pin, _type(pin), _position(pin, _type(pin)), _group(pin), symbol_name=pin.name) for pin in pins]


def assign_unique_symbol_names(pins: list[OrCadPin]) -> None:
    groups: dict[str, list[OrCadPin]] = defaultdict(list)
    for pin in pins:
        pin.symbol_name = pin.source.name
        if pin.source.name:
            groups[pin.source.name.casefold()].append(pin)
    for values in groups.values():
        if len(values) < 2:
            continue
        non_power = [pin for pin in values if pin.electrical_type != 'Power']
        non_power.sort(key=lambda pin: tuple((0, int(part)) if part.isdigit() else (1, part.casefold()) for part in re.split(r'(\d+)', pin.source.number) if part))
        for index, pin in enumerate(non_power, start=1):
            pin.symbol_name = f'{pin.source.name}{index}'


def ambiguous_pins(pins: list[OrCadPin]) -> list[OrCadPin]:
    """Only these pins are eligible for an optional third-party AI suggestion."""
    return [pin for pin in pins if pin.group == 'CONTROL_ANALOG']


def apply_ai_suggestions(pins: list[OrCadPin], suggestions: dict[str, dict[str, str]]) -> None:
    allowed = {pin.source.number: pin for pin in ambiguous_pins(pins)}
    for number, suggestion in suggestions.items():
        pin = allowed.get(number)
        if not pin:
            continue
        electrical_type = suggestion.get('electrical_type')
        position = suggestion.get('position')
        if electrical_type in VALID_TYPES:
            pin.electrical_type = electrical_type
        if position in VALID_POSITIONS:
            pin.position = position
        if electrical_type in VALID_TYPES or position in VALID_POSITIONS:
            pin.classification_source = 'AI suggestion'


def assign_sections(pins: list[OrCadPin], max_pins_per_section: int) -> list[tuple[str, str, list[OrCadPin]]]:
    if not 1 <= max_pins_per_section <= 100:
        raise ValueError('每个 Section 的最大引脚数必须在 1 到 100 之间')
    grouped: dict[str, list[OrCadPin]] = defaultdict(list)
    for pin in pins:
        grouped[pin.group].append(pin)
    sections: list[tuple[str, str, list[OrCadPin]]] = []
    for group in sorted(grouped, key=_group_sort_key):
        values = grouped[group]
        for offset in range(0, len(values), max_pins_per_section):
            name = _section_name(len(sections))
            chunk = values[offset:offset + max_pins_per_section]
            for pin in chunk:
                pin.section = name
            sections.append((name, group, chunk))
    return sections


def build_workbook(device: str, pins: list[OrCadPin], sections: list[tuple[str, str, list[OrCadPin]],], max_pins_per_section: int) -> bytes:
    wb = Workbook()
    paste = wb.active
    paste.title = 'Paste_To_OrCAD'
    paste.append(ORCAD_HEADERS)
    for pin in pins:
        # Capture treats 1 as checked Pin Visibility. PinGroup is intentionally empty: it
        # represents swap groups, not FPGA bank membership.
        paste.append([pin.source.number, pin.symbol_name, pin.electrical_type, '1', 'Line', '', pin.position, pin.section])

    summary = wb.create_sheet('Section_Summary')
    summary.append([f'{device} — OrCAD Section Plan'])
    summary.append([])
    summary.append(['Part Name', device])
    summary.append(['Part Ref Prefix', 'U'])
    summary.append(['Part Numbering', 'Alphabetic'])
    summary.append(['No. of Sections', len(sections)])
    summary.append(['Maximum pins / section', max_pins_per_section])
    summary.append([])
    summary.append(['Section', 'Functional group', 'Pins', 'Note'])
    for section, group, values in sections:
        summary.append([section, group, len(values), 'Pin Visibility = 1; PinGroup left blank'])
    summary.append([])
    summary.append(['Paste rows 2 onward from Paste_To_OrCAD into Capture. Do not copy the header.'])

    audit = wb.create_sheet('Source_Audit')
    audit.append(ORCAD_HEADERS + ['Memory Byte Group', 'Bank', 'I/O Type', 'Functional Group', 'Classification', 'Original Name'])
    for pin in pins:
        audit.append([
            pin.source.number, pin.symbol_name, pin.electrical_type, '1', 'Line', '', pin.position, pin.section,
            pin.source.memory_byte_group, pin.source.bank, pin.source.io_type, pin.group, pin.classification_source, pin.source.name,
        ])

    header_fill = PatternFill('solid', fgColor='1F4E78')
    for sheet in (paste, summary, audit):
        sheet.freeze_panes = 'A2' if sheet != summary else 'A10'
        for cell in sheet[1]:
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
        for row in sheet.iter_rows():
            for cell in row:
                cell.number_format = '@'
                font = copy(cell.font)
                font.name = 'Arial'
                font.sz = 10
                cell.font = font
    for sheet, widths in ((paste, [13, 42, 16, 16, 12, 14, 14, 11]), (summary, [20, 25, 12, 50]), (audit, [13, 42, 16, 16, 12, 14, 14, 11, 18, 10, 12, 20, 18, 42])):
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def convert_ascii_pinout(raw: bytes, max_pins_per_section: int = 100, ai_suggester: Callable[[list[OrCadPin]], dict[str, dict[str, str]]] | None = None) -> tuple[str, bytes, list[OrCadPin], list[tuple[str, str, list[OrCadPin]]]]:
    device, source_pins = parse_ascii_pinout(raw)
    pins = make_orcad_pins(source_pins)
    if ai_suggester:
        apply_ai_suggestions(pins, ai_suggester(ambiguous_pins(pins)))
    assign_unique_symbol_names(pins)
    sections = assign_sections(pins, max_pins_per_section)
    return device, build_workbook(device, pins, sections, max_pins_per_section), pins, sections
